"""A sheet as a plain grid, with merged cells resolved.

Every cell inside a merged range reports the merged value and the range it
belongs to, so a lab merged over three periods reads as three filled cells
that share one anchor.
"""
from __future__ import annotations

from datetime import datetime, time

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

Span = tuple[int, int, int, int]  # min_row, min_col, max_row, max_col


def _to_text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        v = v.time()
    if isinstance(v, time):
        return v.strftime("%H:%M")
    s = str(v).replace("\r", "").strip()
    return s or None


class Grid:
    def __init__(self, values: dict[tuple[int, int], str], spans: dict[tuple[int, int], Span],
                 max_row: int, max_col: int, title: str, transposed: bool = False):
        self._values = values
        self._spans = spans
        self.max_row = max_row
        self.max_col = max_col
        self.title = title
        self.transposed = transposed

    @classmethod
    def from_worksheet(cls, ws: Worksheet) -> Grid:
        values: dict[tuple[int, int], str] = {}
        for row in ws.iter_rows():
            for c in row:
                t = _to_text(c.value)
                if t is not None:
                    values[(c.row, c.column)] = t
        spans: dict[tuple[int, int], Span] = {}
        for rng in ws.merged_cells.ranges:
            span = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    spans[(r, c)] = span
        return cls(values, spans, ws.max_row, ws.max_column, ws.title)

    def transpose(self) -> Grid:
        """Swap rows and columns, for timetables with days across the top."""
        values = {(c, r): v for (r, c), v in self._values.items()}
        spans = {(c, r): (s[1], s[0], s[3], s[2]) for (r, c), s in self._spans.items()}
        return Grid(values, spans, self.max_col, self.max_row, self.title, not self.transposed)

    # ------------------------------------------------------------ access

    def span(self, r: int, c: int) -> Span:
        return self._spans.get((r, c), (r, c, r, c))

    def value(self, r: int, c: int) -> str | None:
        """Cell value, filled in from the merge anchor for merged cells."""
        r0, c0, _, _ = self.span(r, c)
        return self._values.get((r0, c0))

    def row(self, r: int) -> list[str | None]:
        return [self.value(r, c) for c in range(1, self.max_col + 1)]

    def row_text(self, r: int) -> str:
        seen, parts = set(), []
        for c in range(1, self.max_col + 1):
            anchor = self.span(r, c)[:2]
            v = self.value(r, c)
            if v and anchor not in seen:
                seen.add(anchor)
                parts.append(v)
        return " | ".join(parts)

    def cells(self):
        """Anchor cells only: (row, col, value)."""
        for (r, c), v in sorted(self._values.items()):
            yield r, c, v

    def coord(self, r: int, c: int) -> str:
        """Excel-style address in the original sheet, even when transposed."""
        if self.transposed:
            r, c = c, r
        return f"{get_column_letter(c)}{r}"

    def span_coord(self, span: Span) -> str:
        r0, c0, r1, c1 = span
        a, b = self.coord(r0, c0), self.coord(r1, c1)
        return a if a == b else f"{a}:{b}"

    def preview(self, max_rows: int = 45, max_len: int = 40) -> str:
        """Compact text dump for an LLM: one line per non-empty row with coordinates."""
        lines = []
        for r in range(1, min(self.max_row, max_rows) + 1):
            cells = []
            for c in range(1, self.max_col + 1):
                if self.span(r, c)[:2] != (r, c):
                    continue
                v = self._values.get((r, c))
                if v:
                    s = self.span(r, c)
                    tag = self.span_coord(s)
                    cells.append(f"{tag}={v.replace(chr(10), ' ')[:max_len]!r}")
            if cells:
                lines.append(f"row {r}: " + ", ".join(cells))
        return "\n".join(lines)
