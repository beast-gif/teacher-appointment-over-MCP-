"""Step 2: read the timetable and print what we understood.

Nothing is saved yet. Run:
    python -m app.ingestion.reader tests/fixtures/eac_7th_sem.xlsx
"""
from __future__ import annotations

import re
import sys

import openpyxl

from .grid import Grid
from .text import (
    DAY_LABELS,
    find_default_room,
    find_section_info,
    is_break,
    parse_day,
    parse_ltpc,
    parse_time_range,
    split_faculty,
    split_title_rooms,
)


def find_days(grid: Grid) -> tuple[int, dict[int, int]]:
    """Column holding day names, and {row: weekday}."""
    by_col: dict[int, dict[int, int]] = {}
    for r, c, v in grid.cells():
        d = parse_day(v)
        if d is not None:
            by_col.setdefault(c, {})[r] = d
    if not by_col:
        raise ValueError("no day names found")
    day_col = max(by_col, key=lambda c: len(by_col[c]))
    return day_col, by_col[day_col]


def find_time_header(grid: Grid, day_col: int, first_day_row: int) -> int:
    """Row above the days with the most time ranges like '8:10-9:00'."""
    best, best_count = None, 0
    for r in range(1, first_day_row):
        n = sum(1 for c in range(day_col + 1, grid.max_col + 1) if parse_time_range(grid.value(r, c)))
        if n > best_count:
            best, best_count = r, n
    if best is None:
        raise ValueError("no time header row found")
    return best


def read_slots(grid: Grid, header_row: int, day_col: int):
    """Class periods and breaks from the header row."""
    slots, breaks = [], []
    for c in range(day_col + 1, grid.max_col + 1):
        v = grid.value(header_row, c)
        if is_break(v):
            breaks.append({"col": c, "label": " ".join(v.split())})
        elif t := parse_time_range(v):
            slots.append({"col": c, "start": t[0], "end": t[1]})
    # Break times = gap between the periods on either side
    for b in breaks:
        before = [s for s in slots if s["col"] < b["col"]]
        after = [s for s in slots if s["col"] > b["col"]]
        if before and after:
            b["start"], b["end"] = before[-1]["end"], after[0]["start"]
    return slots, breaks


def read_lab_slots(grid: Grid, header_row: int, first_day_row: int, day_col: int):
    """Rows between the header and Monday labelled 'Lab ...', e.g. 'Lab Slots'."""
    labs = []
    for r in range(header_row + 1, first_day_row):
        if not re.search(r"lab", grid.value(r, day_col) or "", re.I):
            continue
        seen = set()
        for c in range(day_col + 1, grid.max_col + 1):
            span = grid.span(r, c)
            if span in seen:
                continue
            seen.add(span)
            if t := parse_time_range(grid.value(r, c)):
                labs.append({"min_col": span[1], "max_col": span[3], "start": t[0], "end": t[1],
                             "cell": grid.span_coord(span)})
    return labs


def find_legend(grid: Grid) -> dict | None:
    """Header row with both a 'code' column and a 'faculty' column."""
    for r in range(1, grid.max_row + 1):
        cols = {c: (grid.value(r, c) or "").lower() for c in range(1, grid.max_col + 1)}
        code = next((c for c, v in cols.items() if "code" in v), None)
        fac = next((c for c, v in cols.items() if re.search(r"faculty|staff|teacher|instructor", v)), None)
        if code and fac and code != fac:
            title = next((c for c, v in cols.items() if c not in (code, fac)
                          and re.search(r"title|subject|course", v)), None)
            ltpc = next((c for c, v in cols.items() if re.search(r"l\s*t\s*p|credit", v)), None)
            return {"header_row": r, "code_col": code, "title_col": title, "faculty_col": fac, "ltpc_col": ltpc}
    return None


def read_legend(grid: Grid, legend: dict) -> list[dict]:
    courses = []
    for r in range(legend["header_row"] + 1, grid.max_row + 1):
        code = grid.value(r, legend["code_col"])
        if not code:
            if courses:
                break  # blank row ends the table
            continue
        title, rooms = split_title_rooms(grid.value(r, legend["title_col"]) or "") if legend["title_col"] else ("", [])
        ltpc, credits, is_lab = parse_ltpc(grid.value(r, legend["ltpc_col"])) if legend["ltpc_col"] else (None, None, False)
        courses.append({
            "code": re.sub(r"\s+", "", code).upper(),
            "title": title,
            "rooms": rooms,
            "ltpc": ltpc,
            "credits": credits,
            "is_lab": is_lab,
            "faculty": split_faculty(grid.value(r, legend["faculty_col"])),
            "row": r,
        })
    return courses


def read_header_info(grid: Grid, header_row: int) -> dict:
    text = "\n".join(grid.row_text(r) for r in range(1, header_row))
    info = find_section_info(text)
    info["default_room"] = find_default_room(text)
    return info


def main(path: str) -> None:
    ws = openpyxl.load_workbook(path).worksheets[0]
    grid = Grid.from_worksheet(ws)

    day_col, day_rows = find_days(grid)
    header_row = find_time_header(grid, day_col, min(day_rows))
    slots, breaks = read_slots(grid, header_row, day_col)
    labs = read_lab_slots(grid, header_row, min(day_rows), day_col)
    info = read_header_info(grid, header_row)
    legend = find_legend(grid)
    courses = read_legend(grid, legend) if legend else []

    fmt = lambda t: t.strftime("%H:%M")
    print(f"Sheet: {ws.title}")
    print(f"Section: {info}\n")

    print(f"Days (column {day_col}):")
    for r, d in day_rows.items():
        print(f"  row {r}: {DAY_LABELS[d]}")

    print(f"\nPeriods (header row {header_row}):")
    for s in slots:
        print(f"  {grid.coord(header_row, s['col'])}: {fmt(s['start'])}-{fmt(s['end'])}")

    print("\nBreaks:")
    for b in breaks:
        times = f"{fmt(b['start'])}-{fmt(b['end'])}" if "start" in b else "?"
        print(f"  {b['label']}: {times}")

    print("\nLab slots:")
    for l in labs:
        print(f"  {l['cell']}: {fmt(l['start'])}-{fmt(l['end'])}")

    print(f"\nLegend: {len(courses)} courses (header row {legend['header_row'] if legend else '-'})")
    for c in courses:
        flag = "  <- lab" if c["is_lab"] else ""
        print(f"  {c['code']}  {c['title']}  rooms={c['rooms'] or '-'}  credits={c['credits']}{flag}")
        for f in c["faculty"]:
            warn = "   <- name cut off" if f.rstrip().endswith("...") else ""
            print(f"      - {f}{warn}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures/eac_7th_sem.xlsx")
